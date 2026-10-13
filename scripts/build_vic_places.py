"""Suburbs, postcodes and council areas -> Victorian state districts, from ABS
allocation files, into scripts/hubs/vic-places-2025.json for the election
hub's "Which seat am I in?" finder.

Each place is split across districts by its usual residents in the 2021
Census, joining the two by mesh block. All inputs are ABS, CC BY 4.0:
  ASGS Edition 3 allocation files SED_2025, SAL_2021, POA_2021 and LGA_2025;
  Census of Population and Housing: Mesh Block Counts, 2021.
SED 2025 follows the VEC districts in force since 1 November 2022. Postal
areas (POA) are the ABS's approximation of delivery postcodes.

    python3 -I scripts/build_vic_places.py <dir holding the five .xlsx files>
"""
import hashlib, json, re, sys
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ABS = 'https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/access-and-downloads/allocation-files/'
FILES = {
    'SED_2025_AUST.xlsx': ABS + 'SED_2025_AUST.xlsx', 'SAL_2021_AUST.xlsx': ABS + 'SAL_2021_AUST.xlsx',
    'POA_2021_AUST.xlsx': ABS + 'POA_2021_AUST.xlsx', 'LGA_2025_AUST.xlsx': ABS + 'LGA_2025_AUST.xlsx',
    'MB_COUNTS_2021.xlsx': 'https://www.abs.gov.au/census/guide-census-data/mesh-block-counts/2021/Mesh%20Block%20Counts%2C%202021.xlsx',
}
HERE = Path(__file__).parent
OUT = HERE / 'hubs' / 'vic-places-2025.json'


def rows(path, sheet=None):
    book = load_workbook(path, read_only=True)
    for ws in ([book[sheet]] if sheet else book.worksheets[:1]):
        it = ws.iter_rows(values_only=True)
        head = next(it)
        for r in it:
            if r and r[0]:
                yield dict(zip(head, r))


def counts(path):
    """Usual residents per Victorian mesh block (the Victoria tables)."""
    book = load_workbook(path, read_only=True)
    out = {}
    for ws in book.worksheets:
        if not ws.title.startswith('Table'):
            continue
        it = ws.iter_rows(values_only=True)
        for r in it:
            if r and r[0] == 'MB_CODE_2021':
                break
        for r in it:
            if r and r[0] and str(r[0]).startswith('2') and len(str(r[0])) == 11 and r[5] == 2:
                out[str(r[0])] = (int(r[4] or 0), int(r[3] or 0), float(r[2] or 0))
    return out


def label(name):
    # ABS qualifies names shared with other places: "Richmond (Vic.)",
    # "Bellfield (Grampians - Vic.)". Keep only the part that tells Victorians apart.
    name = re.sub(r'\s*\(Vic\.\)$', '', name)
    return re.sub(r'\s*\((.+?) - Vic\.\)$', r' (\1)', name)


def main():
    src = Path(sys.argv[1])
    config = json.loads((HERE / 'hubs' / 'vic-election-2026.json').read_text())
    regions = {d['name']: d['region'] for d in config['districts']}
    districts = set(regions)
    # ABS names a Victorian SED with its region: "Albert Park (Southern Metropolitan)".
    sed, seen = {}, set()
    for r in rows(src / 'SED_2025_AUST.xlsx'):
        if r['STATE_CODE_2021'] != '2':
            continue
        m = re.fullmatch(r'(.+) \((.+)\)', r['SED_NAME_2025'])
        if m and m.group(1) in districts:
            if regions[m.group(1)] != m.group(2):
                sys.exit(f'SED 2025 puts {m.group(1)} in {m.group(2)}, not the VEC region {regions[m.group(1)]}')
            sed[r['MB_CODE_2021']] = m.group(1)
            seen.add(m.group(1))
    if seen != districts:
        sys.exit(f'SED 2025 names differ from the VEC list: missing {sorted(districts - seen)}')
    people = counts(src / 'MB_COUNTS_2021.xlsx')
    places = {
        'suburb': {r['MB_CODE_2021']: label(r['SAL_NAME_2021']) for r in rows(src / 'SAL_2021_AUST.xlsx') if r['STATE_CODE_2021'] == '2'},
        'postcode': {r['MB_CODE_2021']: r['POA_CODE_2021'] for r in rows(src / 'POA_2021_AUST.xlsx') if r['MB_CODE_2021'].startswith('2')},
        'council': {r['MB_CODE_2021']: r['LGA_NAME_2025'] for r in rows(src / 'LGA_2025_AUST.xlsx') if r['STATE_CODE_2021'] == '2'},
    }
    order = sorted(districts)
    out = {}
    for kind, mapping in places.items():
        weights = defaultdict(lambda: defaultdict(lambda: [0, 0, 0.0]))
        for mb, place in mapping.items():
            district = sed.get(mb)
            if district not in districts:
                continue  # no usual address, migratory and offshore codes
            p, d, a = people.get(mb, (0, 0, 0.0))
            w = weights[place][district]
            w[0] += p; w[1] += d; w[2] += a
        rows_out = []
        for place, split in weights.items():
            # Residents decide the split; places with none fall back to dwellings, then area.
            col = 0 if sum(v[0] for v in split.values()) else 1 if sum(v[1] for v in split.values()) else 2
            total = sum(v[col] for v in split.values())
            parts = sorted(((order.index(k), v[col] / total) for k, v in split.items() if v[col] > 0), key=lambda x: -x[1])
            if not parts:
                continue
            # Shares in tenths of a percent; any district with residents is kept, however small.
            rows_out.append([place, [[i, max(1, round(s * 1000))] for i, s in parts], sum(v[0] for v in split.values())])
        out[kind] = sorted(rows_out, key=lambda r: r[0])
    digest = {name: hashlib.sha256((src / name).read_bytes()).hexdigest() for name in FILES}
    OUT.write_text(json.dumps({
        'source': 'Australian Bureau of Statistics', 'licence': 'CC BY 4.0',
        'attribution': 'Source: Australian Bureau of Statistics, ASGS Edition 3 allocation files and 2021 Census mesh block counts (CC BY 4.0).',
        'note': 'Places are split by usual residents at the 2021 Census. Postcodes are ABS postal areas, an approximation of delivery postcodes.',
        'files': FILES, 'sha256': digest, 'districts': order,
        'columns': ['name', '[[district index, share in tenths of a percent]]', 'usual residents 2021'],
        **out,
    }, ensure_ascii=False, separators=(',', ':')) + '\n')
    print({k: len(v) for k, v in out.items()}, OUT.stat().st_size, 'bytes')


if __name__ == '__main__':
    main()
