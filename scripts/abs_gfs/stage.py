#!/usr/bin/env python3
"""Offline XLSX rights review, inventory and staged GFS projection.

Original downloads stay in portal/private/abs-gfs. This module never writes
portal/public, a database, or a remote service. Run from the repository root.
"""
from __future__ import annotations

import hashlib
import json
import datetime
from collections import Counter
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from scripts.abs_gfs.parse import CC, TERMS, RELEASE, HEADLINES, JURISDICTIONS, classification, key, parse_sheet, permitted, reconcile, compare_level_tables, parse_net_debt


def load(path):
    return json.loads(path.read_text())


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n")


def review_workbook(workbook):
    notices, exceptions, sheets, all_years = [], [], [], set()
    for sheet in workbook:
        years, stock_dates = set(), set()
        for row in sheet:
            for cell in row:
                value = cell.value
                if isinstance(value, datetime.datetime):
                    stock_dates.add(value.date().isoformat())
                if not isinstance(value, str):
                    continue
                date_match = re.search(r"(\d{1,2}) (June|July) (\d{4})", value)
                if date_match:
                    stock_dates.add(f"{date_match[3]}-{'06' if date_match[2]=='June' else '07'}-{int(date_match[1]):02d}")
                if re.fullmatch(r"\d{4}-\d{2}", value):
                    years.add(value)
                if re.search(r"©|copyright|creative commons|licen[cs]e|third.party|all rights reserved|permission", value, re.I):
                    notice = {"sheet": sheet.title, "cell": cell.coordinate, "text": value.strip()}
                    notices.append(notice)
                    if re.search(r"third.party|all rights reserved|permission|licen[cs]e", value, re.I) and not re.search(r"creativecommons.org/licenses/by/4.0", value, re.I):
                        exceptions.append(notice)
        for date in stock_dates:
            if date.endswith("-06-30"):
                years.add(f"{int(date[:4])-1}-{date[2:4]}")
        all_years.update(years)
        sheets.append({"sheet": sheet.title, "title": sheet.cell(4, 1).value if sheet.title != "Contents" else "Contents",
                       "reference_years": sorted(years), "stock_dates": sorted(stock_dates),
                       "published_column_headers": [str(c.value).strip() for c in sheet[5] if c.value is not None] if sheet.title != "Contents" else []})
    return {"notices": notices, "exceptions": exceptions, "sheets": sheets, "reference_years": sorted(all_years)}


def annotate_auxiliary(document):
    """Inventory mixed cubes at table grain; original headings remain alongside."""
    three_sectors = ["General government", "Non-financial public sector", "Total public sector"]
    five_levels = ["Commonwealth", "Control n.f.d.", "State", "Local", "All levels of government"]
    for sheet in document["sheets"]:
        if sheet["sheet"] == "Contents":
            continue
        number = int(sheet["sheet"].split("_")[1])
        jurisdictions = ["Australia"]
        if document["id"] == "55120DO001_202425":
            levels = five_levels
            sectors = [three_sectors[(number-1) % 3]]
        elif document["id"] == "55120DO094_202425":
            levels = five_levels if number in (2, 3, 4, 8) else ["All levels of government"]
            sectors = three_sectors if number in (1, 2, 3, 4, 14, 15, 16) else ["Total public sector"] if number in (5, 8) else ["General government"]
            if number == 7: levels = ["Commonwealth"]
            if number == 9:
                levels = ["Commonwealth", "State", "All levels of government"]
                jurisdictions += list(JURISDICTIONS.values())
            if number == 10:
                levels = ["Local"]
                jurisdictions += [j for j in JURISDICTIONS.values() if j != "Australian Capital Territory"]
            if number in (11, 12, 13): levels = [{11: "Commonwealth", 12: "State", 13: "Local"}[number]]
            if number in (14, 15, 16): levels = ["National", "State", "Local", "All levels of government"]
            if number == 17: levels = ["Commonwealth", "State", "Local"]
        else:  # Published education cube, including its separately labelled context.
            sectors = ["General government"]
            levels = ["All levels of government"]
            if number in (2, 3, 4, 5, 6):
                levels = ["Commonwealth", "State and local", "Control n.f.d.", "All levels of government"]
                jurisdictions += list(JURISDICTIONS.values())
            if number in (7, 18): levels = ["Commonwealth"]
            if 8 <= number <= 15:
                levels = ["State and local"]
                if number >= 9: jurisdictions = [JURISDICTIONS[str(number-8)]]
            if number == 16: levels, jurisdictions = ["State"], ["Australian Capital Territory"]
            if number == 17: levels = ["Control n.f.d."]
            if number == 19:
                levels, sectors = ["Not applicable: context table"], ["Private expenditure and GDP context; not a public-sector breakdown"]
        sheet["coverage"] = {"levels": levels, "sectors": sectors, "jurisdictions": jurisdictions}
    for field in ("levels", "sectors", "jurisdictions"):
        document[field] = sorted({v for s in document["sheets"] for v in s.get("coverage", {}).get(field, [])})


def projection(inventory, rows):
    documents = {d["id"] + ".xlsx": d for d in inventory["workbooks"] if permitted(d)}
    public = []
    for row in rows:
        d = documents.get(row.get("workbook"))
        dimensions = ([d["dimensions"]] if d and d.get("dimensions") else d.get("series_dimensions", []) if d else [])
        if (not d or row.get("published") is not True or row.get("measure") not in HEADLINES
                or row.get("source_url") != d["url"] or row.get("sha256") != d["receipt"]["sha256"]
                or row.get("snapshot") != d["receipt"]["snapshot"]
                or not any(all(row.get(k) == v for k, v in x.items()) for x in dimensions)):
            continue
        public.append(row)
    return {"schema": 1, "generated_at": inventory["snapshot"], "release": inventory["release"],
            "release_date": inventory["release_date"], "release_url": RELEASE,
            "licence": {"name": "CC BY 4.0", "url": CC, "terms_url": TERMS,
                        "attribution": "Australian Bureau of Statistics (2024–25), Government Finance Statistics, Annual, accessed 17 October 2026. © Commonwealth of Australia 2026.",
                        "changes": "Selected published cells reformatted as a tidy series. Displayed amounts use m/bn; source values remain AUD million. No government or sector totals calculated by OPAX."},
            "documents": [{k: d[k] for k in ("id", "title", "url", "release", "release_date", "status", "dimensions", "receipt", "licence", "exceptions", "series_dimensions") if k in d}
                          for d in documents.values() if any(r["workbook"] == d["id"] + ".xlsx" for r in public)], "rows": public}


def build(root=ROOT):
    stage = root / "data/staged/abs-gfs"
    inventory = load(stage / "inventory.json")
    ledger = load(stage / "receipts.json")
    receipts = {r["url"]: r for r in ledger if r.get("status") == 200 and not r.get("challenge")}
    for url in (TERMS, "https://www.abs.gov.au/how-cite-abs-sources"):
        if url not in receipts:
            raise ValueError("Missing rights receipt: " + url)
    rows, debt_checks = [], []
    for document in inventory["workbooks"]:
        receipt = receipts.get(document["url"])
        path = root / "portal/private/abs-gfs" / (document["id"] + ".xlsx")
        if not receipt or not path.exists():
            document.update(status="held", hold_reason="Original workbook not acquired; rights unverified")
            continue
        if hashlib.sha256(path.read_bytes()).hexdigest() != receipt["sha256"]:
            raise ValueError("Source hash mismatch: " + document["id"])
        workbook = load_workbook(path, data_only=True)
        review = review_workbook(workbook)
        document.update(review, dimensions=classification(document["title"]),
                        receipt={k: receipt[k] for k in ("snapshot", "sha256", "status", "challenge")},
                        original_path=str(path.relative_to(root)), status="held" if review["exceptions"] else "cleared",
                        licence={"status": "held" if review["exceptions"] else "open", "name": "CC BY 4.0", "url": CC,
                                 "terms_url": TERMS, "terms_sha256": receipts[TERMS]["sha256"], "workbook_reviewed": True,
                                 "basis": "ABS website grant applies to published numerical tables; workbook notices inspected on every sheet. Coat of Arms, ABS logo, trade marks and third-party material excluded."})
        document.pop("hold_reason", None)
        dimensions = document["dimensions"]
        if dimensions:
            document["levels"] = [dimensions["level_of_government"]]
            document["sectors"] = [dimensions["sector"]]
            document["jurisdictions"] = [dimensions["jurisdiction"]]
            for sheet in document["sheets"]:
                if sheet["sheet"] != "Contents":
                    sheet["dimensions"] = dimensions
        else:
            # Auxiliary cubes contain multiple consolidation groups. Retain the
            # original sheet titles/column labels rather than invent a single one.
            document["levels"] = ["Multiple published levels; see sheet titles and column headers"]
            document["sectors"] = ["Multiple published sectors; see sheet titles and column headers"]
            document["jurisdictions"] = ["Multiple published jurisdictions; see sheet titles and column headers"]
            document["scope_note"] = "Auxiliary cube; each sheet is inventoried with original title, column headers and periods. Only the six explicit Table_15 L2 cells from Key Tables enter the core series; education/function tables are not staged."
            annotate_auxiliary(document)
        if review["exceptions"]:
            document["hold_reason"] = "Workbook-specific notice requires review"
        if permitted(document):
            for sheet in workbook:
                rows.extend(parse_sheet(list(sheet.values), sheet.title, document))
            if document["id"] == "55120DO094_202425":
                debt_rows, debt_checks = parse_net_debt(workbook, document)
                rows.extend(debt_rows)
        workbook.close()
    inventory["snapshot"] = max(r.get("completed_at", r["snapshot"]) for r in ledger)
    duplicates = Counter(key(row) for row in rows)
    if any(n != 1 for n in duplicates.values()):
        raise ValueError("Duplicate series key")
    checks = reconcile(rows) + debt_checks
    comparison = next(d for d in inventory["workbooks"] if d["id"] == "55120DO001_202425")
    if permitted(comparison):
        workbook = load_workbook(root / comparison["original_path"], data_only=True)
        checks.extend(compare_level_tables(workbook, comparison, rows))
        workbook.close()
    if any(c["status"] == "fail" for c in checks):
        write(stage / "reconciliation.json", checks)
        raise ValueError("Published identity does not reconcile")
    write(stage / "inventory.json", inventory)
    tables = [{"workbook": d["id"]+".xlsx", "url": d["url"], "release": d["release"], "release_date": d["release_date"], "format": d["format"],
               "licence": d["licence"], "receipt": d.get("receipt"), **sheet,
               "coverage": sheet.get("coverage", {"levels": d["levels"], "sectors": d["sectors"], "jurisdictions": d["jurisdictions"]})}
              for d in inventory["workbooks"] for sheet in d.get("sheets", []) if sheet["sheet"] != "Contents"]
    write(stage / "tables.json", tables)
    write(stage / "reconciliation.json", checks)
    public_rows = [r for r in rows if r["measure"] in HEADLINES]
    (stage / "series.jsonl").write_text("".join(json.dumps(r, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n" for r in public_rows))
    hub = projection(inventory, public_rows)
    write(stage / "hub.json", hub)
    summary = {"release": inventory["release"], "workbooks": len(inventory["workbooks"]),
               "tables": sum(len(d.get("sheets", []))-1 for d in inventory["workbooks"] if d.get("sheets")),
               "inventory_core_by_level": dict(Counter(d["dimensions"]["level_of_government"] for d in inventory["workbooks"] if d.get("dimensions"))),
               "inventory_core_by_sector": dict(Counter(d["dimensions"]["sector"] for d in inventory["workbooks"] if d.get("dimensions"))),
               "auxiliary_workbooks": [d["id"] for d in inventory["workbooks"] if not d.get("dimensions")],
               "cleared_workbooks": sum(permitted(d) for d in inventory["workbooks"]),
               "held_workbooks": sum(not permitted(d) for d in inventory["workbooks"]),
               "rows": len(public_rows), "exported_rows": len(hub["rows"]),
               "by_level": dict(Counter(r["level_of_government"] for r in public_rows)),
               "by_sector": dict(Counter(r["sector"] for r in public_rows)),
               "measures": dict(Counter(r["measure"] for r in public_rows)),
               "reconciliation": dict(Counter(c["status"] for c in checks)),
               "years": sorted({r["reference_year"] for r in public_rows})}
    write(stage / "summary.json", summary)
    print(json.dumps(summary, indent=2))
    return hub


if __name__ == "__main__":
    build()
