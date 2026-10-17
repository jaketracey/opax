#!/usr/bin/env python3
"""Offline statement review and staging. Never writes public assets or databases."""
from __future__ import annotations
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from scripts.tax_expenditures.parse import CC, TERMS, STATEMENT, SHEET, TOTALS, CATEGORIES, NS, document_metadata, tables, text, key, permitted, parse_sheet, parse_combined


def load(path):
    return json.loads(path.read_bytes())


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(',', ':'))+'\n')


def projection(inventory, rows):
    docs = {d['id']:d for d in inventory['documents'] if permitted(d)}
    items = {i['item_id']:i for i in inventory['items']}
    cleared = []
    for r in rows:
        d, i = docs.get(r.get('source_file')), items.get(r.get('item_id'))
        if not d or not i or r.get('published') is not True or r.get('estimate_basis') != 'revenue forgone':
            continue
        if (r.get('source_url') != d['url'] or r.get('sha256') != d['receipt']['sha256'] or r.get('snapshot') != d['receipt']['snapshot']
            or r.get('statement') != d['statement'] or r.get('licence_url') != CC or r.get('unit') != 'AUD million'
            or any(r.get(k) != i[k] for k in ('title','category','benchmark','record_kind','document_table','reliability','tax_expenditure_type'))):
            continue
        if i['record_kind']=='item' and (d['format']!='XLSX' or r.get('sheet') != SHEET or r.get('cell') != i['cells'].get(r.get('measure_year'))):
            continue
        if i['record_kind']=='published combined estimate' and (d['format']!='DOCX' or r.get('sheet')!=i['table'] or r.get('cell') != i['cells'].get(r.get('measure_year'))):
            continue
        cleared.append(r)
    return {'schema':1,'generated_at':inventory['snapshot'],'statement':inventory['statement'],
        'statement_url':inventory['statement_url'],'published_date':inventory['published_date'],
        'licence':{'name':'CC BY 4.0','url':CC,'terms_url':TERMS,
            'attribution':'Based on Commonwealth of Australia data. © Commonwealth of Australia 2025. Treasury, 2025–26 Tax Expenditures and Insights Statement.',
            'changes':'Published cells selected and reformatted; displayed in m/bn with original AUD million values retained. No item, category or year estimates added by OPAX.'},
        'items':list(items.values()),'documents':[d for d in docs.values() if any(r['source_file']==d['id'] for r in cleared)],'rows':cleared}


def build(root=ROOT):
    stage, cache = root/'data/staged/tax-expenditures', root/'portal/private/tax-expenditures'
    inventory, receipts = load(stage/'inventory.json'), load(stage/'receipts.json')
    ledger = {r['url']:r for r in receipts}
    if any(r.get('held') for r in receipts):
        raise ValueError('Source access is held')
    terms = ledger[TERMS]
    if terms['status']!=200 or terms['challenge'] or 'Creative Commons Attribution 4.0' not in (cache/'copyright.html').read_text():
        raise ValueError('Missing applicable website grant')
    docx = next(d for d in inventory['documents'] if d['format']=='DOCX')
    pdf = next(d for d in inventory['documents'] if d['format']=='PDF')
    xlsx = next(d for d in inventory['documents'] if d['format']=='XLSX')
    # Hashes and rights are checked before numerical parsing.
    for d in inventory['documents']:
        receipt = ledger.get(d['url'], {})
        path = cache/d['id']
        if receipt.get('status')!=200 or receipt.get('challenge') or hashlib.sha256(path.read_bytes()).hexdigest()!=receipt['sha256']:
            raise ValueError('Original hash/access mismatch: '+d['id'])
        d['receipt'] = {k:receipt[k] for k in ('status','challenge','held','snapshot','sha256')}
        d['original_path'] = str(path.relative_to(root))
    xml, nodes, grids = tables(cache/docx['id'])
    paragraphs = [text(p) for p in xml.findall('.//w:p',NS)]
    grant = next(p for p in paragraphs if p.startswith('This publication is available for your use under Creative Commons Attribution 4.0'))
    workbook = load_workbook(cache/xlsx['id'],data_only=True)
    notices=[]
    for sheet in workbook:
        for row in sheet:
            for c in row:
                if isinstance(c.value,str) and re.search(r'©|copyright|creative commons|all rights reserved|permission to reproduce|third.party (material|content)',c.value,re.I):
                    notices.append({'sheet':sheet.title,'cell':c.coordinate,'text':c.value})
    if notices:
        raise ValueError('Workbook rights notice requires a new review')
    for d in inventory['documents']:
        d.update(status='cleared',exceptions=[],licence={'status':'open','name':'CC BY 4.0','url':CC,'terms_url':TERMS,
            'terms_sha256':terms['sha256'],'document_reviewed':True,'document_grant_sha256':docx['receipt']['sha256'],
            'scope':'Treasury-produced statement and its published numerical tables. Coat of Arms, Treasury logo, third-party and otherwise marked material excluded.',
            'attribution':'Based on Commonwealth of Australia data. © Commonwealth of Australia 2025.'})
    inventory['rights_review']={'document_grant':grant,'document_grant_locator':'DOCX front matter; PDF page 2',
        'workbook_notices':notices,'review':'All 35 worksheets and statement rights notices checked; no narrower notice on staged item or combined tables. Branding and narrative/distributional chart content not projected.'}
    items,combined=document_metadata(cache/docx['id'])
    if set(combined)!=set(TOTALS):
        raise ValueError('Missing published combined tables')
    sheet=workbook[SHEET];grid=list(sheet.values)
    years=grid[0][2:]
    for row_index,(code,title,*values) in enumerate(grid[1:],2):
        items[code]['title_as_statement']=items[code]['title'];items[code]['title']=title
        items[code]['cells']={year:chr(67+i)+str(row_index) for i,year in enumerate(years)}
    for i in combined.values():
        i['cells']={year:f'R2C{n}' for n,year in enumerate(i['years'],1)}
    rows=parse_sheet([list(r) for r in grid],xlsx,items)+parse_combined(docx,combined)
    checks=[]
    for row in rows:
        if row['record_kind']!='item':continue
        item=items[row['item_id']];index=item['years'].index(row['measure_year'])
        actual=item['values'][0 if item.get('included_in') else index];normalized=re.sub(r'[\s,]', '',actual)
        expected=row['source_value']
        same=(expected=='-' and actual=='Included in '+item['included_in']) if item.get('included_in') else normalized==str(expected)
        checks.append({'check':'Workbook nil / statement inclusion reference' if item.get('included_in') else 'XLSX versus statement item cell','item_id':row['item_id'],'year':row['measure_year'],
            'workbook_cell':row['cell'],'document_table':item['document_table'],'status':'pass' if same else 'fail',
            'workbook_value':expected,'statement_value':actual})
    pdf_text=subprocess.run(['pdftotext','-layout',str(cache/pdf['id']),'-'],capture_output=True,text=True,check=True).stdout
    pages=pdf_text.split('\f')
    for identity,item in combined.items():
        page=pages[item['pdf_page']-1]
        match=re.search(re.escape(item['table'])+r' Combined estimate of aggregate revenue forgone \(\$m\)\s*\n([^\n]+)\n([^\n]+)',page)
        if not match or match[1].split()!=item['years']:raise ValueError('Combined PDF control table missing')
        controls=match[2].split()
        for row in (r for r in rows if r['item_id']==identity):
            control=controls[item['years'].index(row['measure_year'])]
            checks.append({'check':'Published combined total: DOCX versus PDF','item_id':identity,'year':row['measure_year'],
                'table':item['table'],'pdf_page':item['pdf_page'],'value_million':row['value'],'publisher_value':control,
                'status':'pass' if int(control.replace(',',''))==row['value'] else 'fail'})
    count_grid=next(g for g in grids if g and g[0]==['Tax expenditure category','Number of tax expenditures'])
    counts=Counter(r['item_id'][0] for r in rows if r['record_kind']=='item' and r['measure_year']==inventory['statement'])
    for letter,category in CATEGORIES.items():
        expected=int(next(r[1] for r in count_grid[1:] if r[0]==category))
        checks.append({'check':'Published item count','category':category,'count':counts[letter],'publisher_count':expected,'status':'pass' if counts[letter]==expected else 'fail'})
    expected=int(next(r[1] for r in count_grid[1:] if r[0]=='Total'))
    checks.append({'check':'Published item count','category':'All items','count':len(items),'publisher_count':expected,'status':'pass' if len(items)==expected else 'fail'})
    if any(c['status']!='pass' for c in checks):
        write(stage/'reconciliation.json',checks)
        raise ValueError('Statement reconciliation failed; inspect reconciliation.json')
    inventory['snapshot']=max(r['completed_at'] for r in receipts)
    inventory['items']=[{k:v for k,v in i.items() if k not in ('years','values')} for i in [*items.values(),*combined.values()]]
    inventory['sheets']=[{'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'staged':s.title==SHEET} for s in workbook]
    inventory['scope']='Latest statement only. Earlier release vintages are not acquired or staged; revised historical years belong to this statement.'
    inventory['held']=[]
    inventory['not_staged']=['34 distributional/chart worksheets; outside item-estimate scope','Earlier statement vintages; machine-readable layout and per-file rights not reviewed','Outlay-equivalent estimates: discussed as an alternative method, not published as a numerical series in this statement']
    workbook.close()
    if len({key(r) for r in rows})!=len(rows):raise ValueError('Duplicate observation')
    write(stage/'inventory.json',inventory);write(stage/'reconciliation.json',checks)
    (stage/'series.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False,allow_nan=False,separators=(',',':'))+'\n' for r in rows))
    hub=projection(inventory,rows);write(stage/'hub.json',hub)
    summary={'statement':inventory['statement'],'items':len(items),'published_combined_groups':len(combined),'item_rows':sum(r['record_kind']=='item' for r in rows),
        'combined_rows':sum(r['record_kind']!='item' for r in rows),'rows':len(rows),'years':sorted({r['measure_year'] for r in rows}),
        'negative_estimates':sum(r['value'] is not None and r['value']<0 for r in rows),'estimate_status':dict(Counter(r['estimate_status'] for r in rows)),
        'reconciliation':dict(Counter(c['status'] for c in checks)),'held_documents':0,'outlay_equivalent_rows':0}
    write(stage/'summary.json',summary);print(json.dumps(summary,indent=2));return hub

if __name__=='__main__':build()
