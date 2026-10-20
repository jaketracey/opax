"""Offline staging only. Does not promote, fetch, write public assets or approve."""
from collections import Counter
from datetime import datetime
import hashlib
import json
import re

from openpyxl import load_workbook

from .acquire import STAGE, CACHE, BASE, DATASETS
from .inventory import resources
from .parse import budget_cells, function_cells, function_code

CC = 'https://creativecommons.org/licenses/by/4.0/'
SELECTED = ('1.3', 'A.1', 'A.4', 'A.7')
FUNCTION = '3f756590-9e67-4a05-97a0-f594247d2c72'
WARNING = 'This resources provides the Final Budget Outcome for Government sector expenses categorised by the 4 digit Government Purpose Classification and the Economic Type Framework ($million). This data is more detailed than reported by Department of Treasury and Finance to the ABS or the Final Budget Outcome. It should, however, not be relied on and is provided for interest and is subject to revision.'


def write(name, data):
    (STAGE / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n')


def build():
    receipts = json.loads((STAGE / 'receipts.json').read_text())
    datasets = []
    all_resources = []
    for slug in DATASETS:
        url = BASE + '/data/dataset/' + slug
        r = next(r for r in receipts if r['url'] == url)
        html = (CACHE / r['sha256']).read_text()
        assert 'creativecommons.org/licenses/by/4.0' in html
        rows = resources(html)
        dataset = dict(id=slug, url=url, title='Budget 2015-16 Tables and Data' if slug == DATASETS[0] else 'South Australian State Government Finance Statistics',
                       receipt=r, licence='CC BY 4.0', licence_url=CC, resource_count=len(rows))
        datasets.append(dataset)
        for row in rows:
            row.update(dataset=slug, dataset_url=url, dataset_snapshot=r['snapshot'], dataset_sha256=r['sha256'],
                       vintage='2015-16 Budget' if slug == DATASETS[0] else 'Final Budget Outcome; resource last updated 17 March 2017',
                       selected=row['table_number'] in SELECTED if slug == DATASETS[0] else row['id'] == FUNCTION,
                       status='not-selected', row_count=None, sha256=None,
                       licence=dict(name='CC BY 4.0', url=CC, dataset_explicit=True, resource_explicit=False, file_notices_reviewed=False))
            all_resources.append(row)
    assert len([r for r in all_resources if r['dataset'] == DATASETS[0]]) == 117
    series, documents, workbooks = [], [], {}
    for resource in all_resources:
        if not resource['selected']:
            continue
        resource_page = next(r for r in receipts if r['kind'] == 'resource' and r['resource_id'] == resource['id'])
        html = (CACHE / resource_page['sha256']).read_text()
        assert 'creativecommons.org/licenses/by/4.0' in html
        resource['licence'].update(resource_explicit=True, resource_page_sha256=resource_page['sha256'])
        resource['resource_receipt'] = resource_page
        original = re.search(r'<meta name="DCTERMS.Identifier" content="([^"]+)"', html)[1]
        resource['download_url_sha256'] = hashlib.sha256(original.encode()).hexdigest()
        resource['download_url'] = original if resource['id'] != FUNCTION else None
        fetched = next((r for r in receipts if r['kind'] == 'download' and r['resource_id'] == resource['id']), None)
        resource['download_receipt'] = fetched
        if not fetched or not fetched['body_retained']:
            resource.update(status='held', hold='Original workbook unavailable; file notices unread')
            continue
        path = CACHE / fetched['sha256']
        assert hashlib.sha256(path.read_bytes()).hexdigest() == fetched['sha256']
        with path.open('rb') as f:
            book = load_workbook(f, data_only=True)
        # All sheets, hidden states, text, comments and header/footer notices
        # were reviewed in this fixed snapshot. No separate rights sheet or
        # narrower rights notice exists. Store only published notes, not file
        # core-properties authors or other local account metadata.
        notes = [str(c.value) for s in book for row in s for c in row if isinstance(c.value, str) and
                 (c.value.startswith('Note:') or re.match(r'^\([ab]\)', c.value))]
        if resource['table_number'] == 'A.4':
            notes.append(book.active['A52'].value)
        resource['licence'].update(file_notices_reviewed=True, exceptions=[],
                                  reviewed_sheets=[dict(sheet=s.title, state=s.sheet_state) for s in book],
                                  notice_finding='All sheet text, comments and headers/footers reviewed; no narrower rights notice or separate copyright sheet')
        hold = 'Workbook scale unlabelled and inconsistent with catalogue $million; no currency conversion inferred' if resource['id'] == FUNCTION else None
        resource.update(status='held' if hold else 'reviewed-staged', hold=hold, sha256=fetched['sha256'], snapshot=fetched['snapshot'],
                        notes=notes, warnings=[WARNING] if hold else [], workbook=resource['id'] + '.xlsx')
        resource['attribution'] = f"Department of Treasury and Finance, Government of South Australia, {next(d['title'] for d in datasets if d['id'] == resource['dataset'])}, Sourced on {fetched['snapshot'][:10]}, {resource['dataset_url']}"
        rows = [r for sheet in book for r in (function_cells(sheet) if resource['id'] == FUNCTION else budget_cells(sheet, resource['table_number']))]
        resource['row_count'] = len(rows)
        resource['fiscal_years'] = sorted({r['fiscal_year'] for r in rows})
        for row in rows:
            row.update(resource_id=resource['id'], dataset=resource['dataset'], dataset_url=resource['dataset_url'],
                       source_url=resource['resource_url'], workbook=resource['workbook'], table_number=resource['table_number'],
                       vintage=resource['vintage'], snapshot=resource['snapshot'], sha256=resource['sha256'],
                       licence='CC BY 4.0', licence_url=CC, published=True,
                       consolidation='General government sector as published; workbook does not separately describe consolidation adjustments',
                       status=resource['status'])
        series.extend(rows)
        documents.append(resource)
        workbooks[resource['table_number'] or 'function'] = book
    write('inventory.json', dict(schema=1, source='Department of Treasury and Finance, Government of South Australia', datasets=datasets, resources=all_resources))
    (STAGE / 'series.jsonl').write_text(''.join(json.dumps(r, ensure_ascii=False, allow_nan=False) + '\n' for r in series))
    checks = reconcile(workbooks)
    write('reconciliation.json', dict(status='pass' if all(c['passed'] for c in checks) else 'failed',
          checks=checks, public_computed_totals=0, arithmetic_scope='Internal diagnostics only; never project these sums, differences or inferred zero values'))
    intervals = []
    for current in receipts[1:]:
        earlier = [r for r in receipts if r['requested_at'] < current['requested_at']]
        prev = max(earlier, key=lambda r: r['snapshot'])
        gap = (datetime.fromisoformat(current['requested_at']) - datetime.fromisoformat(prev['snapshot'])).total_seconds()
        intervals.append(dict(url=current['url'], requested_at=current['requested_at'], completion_gap_seconds=gap, passed=gap>=5.1))
    write('acquisition-audit.json', dict(status='deviation', exceptions=[i for i in intervals if not i['passed']],
          correction='Both overlapping dataset responses were HTTP 200. Exact receipts restored after ledger race. Persistent process lock added before any subsequent request.',
          post_correction_requests_sequential=all(i['passed'] for i in intervals[3:]), intervals=intervals))
    hub = dict(schema=1, generated_at=max(r['snapshot'] for r in receipts), source='Department of Treasury and Finance, Government of South Australia',
               vintage='2015-16 Budget', documents=documents, rows=[r for r in series if r['status']=='reviewed-staged'],
               review_status='pending', publication='Staged only; independent review and explicit annual promoter required')
    write('hub.json', hub)
    print(json.dumps(dict(resources=len(all_resources), rows=len(series), public_candidates=len(hub['rows']), checks=len(checks), failed=sum(not c['passed'] for c in checks), by_resource=dict(Counter(r['resource_id'] for r in series)))))


def reconcile(books):
    checks = []

    def test(name, sheet, target, components, tolerance=None):
        values = [sheet[cell].value for cell, _ in components]
        if any(not isinstance(v, (int, float)) for v in values):
            return  # Never interpret a blank or a published dash as zero.
        total = sheet[target].value
        if not isinstance(total, (int, float)):
            return
        summed = sum(v * sign for v, (_, sign) in zip(values, components))
        bound = tolerance if tolerance is not None else (len(values) + 1) * .5
        delta = total - summed
        checks.append(dict(check=name, sheet=sheet.title, target=target, components=components,
                           published_total=total, internal_component_sum=summed, delta=delta, tolerance=bound,
                           passed=abs(delta)<=bound + 1e-6))

    definitions = {
        'A.1': [(13, [(r,1) for r in range(5,11)]), (28,[(r,1) for r in [17,19,20,21,22,23,24,25]]),
                (32,[(13,1),(28,-1)]),(40,[(32,1),(36,1)]),(44,[(32,1)]),
                (55,[(48,1),(49,-1),(50,-1),(51,1),(52,1)]),(59,[(44,1),(55,-1)])],
        'A.4': [(16,[(r,1) for r in [6,7,8,9,11,12,13]]),(23,[(19,1),(20,1)]),(26,[(16,1),(23,1)]),
                (38,[(r,1) for r in range(29,36)]),(41,[(26,1),(38,-1)]),(44,[(16,1),(38,-1)]),
                (45,[(38,1),(16,-1),(11,1)]),(46,[(29,1),(30,1),(31,1),(6,-1),(7,-1),(8,-1)])],
        'A.7': [(13,[(r,1) for r in range(5,11)]),(23,[(r,1) for r in range(16,21)]),(26,[(13,1),(23,1)]),
                (34,[(30,1),(31,1)]),(48,[(r,1) for r in range(41,46)]),
                (51,[(26,1),(34,1),(37,1),(39,1),(48,1)]),(53,[(26,1)]),(54,[(34,1)]),(59,[(53,1),(54,1),(56,1)])]}
    for table, defs in definitions.items():
        if table not in books:
            continue
        sheet = books[table].active
        for col in 'BCDEF':
            for target, terms in defs:
                test(table, sheet, f'{col}{target}', [(f'{col}{r}',sign) for r,sign in terms])
    if '1.3' in books:
        summary=books['1.3'].active
        for table, pairs in [('A.1',[(6,32),(7,59)]),('A.4',[(14,46),(16,32)])]:
            if table not in books:
                continue
            other=books[table].active
            for a,b in pairs:
                for col, othercol in zip('CDEFG','BCDEF'):
                    value=summary[f'{col}{a}'].value / 1000
                    published=other[f'{othercol}{b}'].value
                    checks.append(dict(check='Table 1.3 cross-table control',sheet=summary.title,target=f'{col}{a}',
                                       compared_table=table,compared_cell=f'{othercol}{b}',delta=value-published,
                                       tolerance=.5,passed=abs(value-published)<=.500001))
    if 'function' in books:
        for sheet in books['function']:
            end=next(c.column for c in sheet[3] if c.value=='Total GPC')
            totalrow=next(c.row for c in sheet['A'] if c.value=='TOTAL')
            coded=[r for r in range(5,totalrow) if re.fullmatch(r'\d{4}',function_code(sheet.cell(r,1)))]
            for row in coded+[totalrow]:
                # Sparse blanks are not zero; compare only the explicitly
                # published numeric components, recording this diagnostic scope.
                cells=[(sheet.cell(row,c).coordinate,1) for c in range(3,end) if isinstance(sheet.cell(row,c).value,(int,float))]
                test('Function published numeric components to published Total GPC (blank semantics unresolved)',sheet,sheet.cell(row,end).coordinate,cells)
            for col in range(3,end+1):
                cells=[(sheet.cell(r,col).coordinate,1) for r in coded if isinstance(sheet.cell(r,col).value,(int,float))]
                test('Function published numeric components to published TOTAL (blank semantics unresolved)',sheet,sheet.cell(totalrow,col).coordinate,cells)
    return checks


if __name__ == '__main__':
    build()
