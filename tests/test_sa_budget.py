"""Synthetic fixtures reproduce official layouts; no people or real amounts."""
import copy
import json
from pathlib import Path
import tempfile
import shutil
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from scripts.sa_budget import acquire as access
from scripts.sa_budget.check import checked_projection, projection
from scripts.sa_budget.inventory import resources
from scripts.sa_budget.parse import budget_cells, function_cells
from scripts.sa_budget.stage import reconcile, FUNCTION
from scripts import promote_hub_data as promoter


def summary():
    s=Workbook().active
    s.title='SOKGGSBI'
    s['A1']='Table 1.3: Summary of key general government sector budget indicators'
    s['B3']='2013–14';s['B4']='Actual'
    s['C2']='2014–15';s['C3']='Estimated';s['C4']='Result'
    s['D3']='2015–16';s['D4']='Budget'
    s['A6']='Net operating balance ($m)'
    for col,value in [('B',-1234500),('C',22000),('D',43500)]:
        s[col+'6']=value;s[col+'6'].number_format=r'#\ ###,;\-#\ ###,;\—'
    s['A10']='Revenue real growth (%)'
    for col in 'BCD':s[col+'10']=1.27;s[col+'10'].number_format='0.0'
    return s


class SABudgetTests(unittest.TestCase):
    def test_split_headers_and_excel_scale_are_not_confused_with_raw_values(self):
        rows=budget_cells(summary(),'1.3')
        self.assertEqual(len(rows),6)
        self.assertEqual([r['fiscal_year'] for r in rows[:3]],['2013-14','2014-15','2015-16'])
        self.assertEqual(rows[0]['estimate_or_actual'],'actual')
        self.assertEqual(rows[1]['status_as_published'],'Estimated Result')
        self.assertEqual(rows[0]['value'],-1235)
        self.assertEqual(rows[0]['raw_value'],-1234500)
        self.assertEqual(rows[2]['value'],44)
        self.assertEqual(rows[3]['value'],1.3)
        self.assertEqual(rows[3]['units'],'percent')

    def test_unrecognised_scale_fails_closed(self):
        s=summary();s['B6'].number_format='General'
        with self.assertRaisesRegex(ValueError,'money format'):budget_cells(s,'1.3')

    def test_same_grants_label_retains_publisher_revenue_and_expense_context(self):
        s=Workbook().active;s['A1']='Table A.1: General government sector operating statement ($million)'
        s['B2']='2015–16\nBudget';s['A4']='Revenue';s['A6']='Grants';s['B6']=8
        s['A16']='Expenses';s['A25']='Grants';s['B25']=3
        rows=budget_cells(s,'A.1')
        self.assertEqual([r['measure'] for r in rows],['Revenue: Grants','Expenses: Grants'])
        self.assertEqual([r['measure_as_published'] for r in rows],['Grants','Grants'])

    def test_matrix_uses_published_etf_headers_and_formatted_four_digit_gpc(self):
        s=Workbook().active;s.title='2012-13';s['A1']='South Australian GPC by ETF - 2012-13'
        s['A4']='GPC';s['B4']='Description';s['C3']=1145;s['C4']='Assumption of liabilities'
        s['D3']=1211;s['D4']='Funded super';s['E3']='Total GPC'
        s['A5']=541;s['A5'].number_format='0000';s['B5']='Community service fixture'
        s['D5']=4;s['E5']=4;s['A6']='TOTAL';s['C6']=0;s['D6']=4;s['E6']=4
        rows=function_cells(s)
        self.assertEqual(len(rows),5)  # Blank C5 is absent, never an inferred zero.
        self.assertEqual(rows[0]['function_code'],'0541')
        self.assertEqual(rows[0]['function_code_raw'],541)
        self.assertEqual(rows[0]['economic_type_code'],'1211')
        self.assertEqual(rows[0]['unit_status'],'unresolved')
        self.assertEqual({r['cell'] for r in rows},{'D5','E5','C6','D6','E6'})

    def test_multiline_cash_label_and_nil_marker_remain_source_cells(self):
        s=Workbook().active;s.title='Sheet 1';s['A1']='Table A.7: General government sector cash flow statement ($million)'
        s['B2']='\n2015–16\nBudget';s['A36']='Net cash flows from investments in financial'
        s['A37']='assets for policy purposes (b)';s['B37']=8;s['A44']='Dividends paid';s['B44']='—'
        rows=budget_cells(s,'A.7')
        self.assertEqual(rows[0]['measure'],'Net cash flows from investments in financial assets for policy purposes (b)')
        self.assertEqual(rows[1]['value'],None)
        self.assertEqual(rows[1]['source_marker'],'—')

    def test_html_titles_uuid_and_published_table_numbers_not_truncated_display(self):
        html='''<li class="resource-item" data-id="synthetic-id"><a class="heading" href="/data/dataset/example/resource/synthetic-id" title="Table A.14(a) General government sector grant revenue">Truncated...</a><ul><li>Explore</li></ul><span class="format-label" data-format="xlsx">XLSX</span></li>'''
        row=resources(html)[0]
        self.assertEqual(row['table_number'],'A.14(a)')
        self.assertEqual(row['format'],'XLSX')
        self.assertNotIn('Truncated',row['title'])

    def test_reconciliation_accepts_published_rounding_and_reports_failure(self):
        w=Workbook();s=w.active
        for col in 'BCDEF':
            for r in range(5,11):s[f'{col}{r}']=10
            s[col+'13']=61
        checks=reconcile({'A.1':w})
        self.assertTrue(all(c['passed'] for c in checks))
        s['B13']=70
        self.assertTrue(any(not c['passed'] for c in reconcile({'A.1':w})))

    def test_offline_real_snapshot_has_447_candidates_and_4774_held_cells(self):
        _,counts=checked_projection(access.STAGE)
        self.assertEqual(counts['staged_cells'],5221)
        self.assertEqual(counts['candidate_cells'],447)
        self.assertEqual(counts['held_cells'],4774)
        self.assertEqual(counts['budget_checks_passed'],120)
        self.assertEqual(counts['function_failed'],5)

    def test_licence_and_held_document_gates(self):
        h=json.loads((access.STAGE/'hub.json').read_text())
        self.assertEqual(len(projection(h)['rows']),447)
        for key in ('dataset_explicit','resource_explicit','file_notices_reviewed'):
            changed=copy.deepcopy(h);changed['documents'][0]['licence'][key]=False
            self.assertEqual(len(projection(changed)['rows']),405)
        h['documents'][0]['hold']='licence exception'
        out=projection(h)
        self.assertEqual(len(out['rows']),405)
        self.assertFalse(any(r['resource_id']==FUNCTION for r in out['rows']))
        self.assertNotIn('internal_component_sum',json.dumps(out))

    def test_refusal_persists_and_repeated_original_never_hits_transport(self):
        with tempfile.TemporaryDirectory(prefix='opax-sa-test-') as tmp:
            root=Path(tmp);url=access.BASE+'/robots.txt'
            for prior, raises in [({'url':url,'status':403,'challenge':False},True),({'url':url,'status':200,'sha256':'cached'},False)]:
                (root/'receipts.json').write_text(json.dumps([prior]))
                with patch.object(access,'STAGE',root),patch.object(access,'CACHE',root/'cache'),patch.object(access.request,'build_opener') as transport:
                    if raises:
                        with self.assertRaisesRegex(ValueError,'Host stopped'):access.acquire(url,'robots')
                    else:self.assertEqual(access.acquire(url,'robots')['sha256'],'cached')
                    transport.assert_not_called()
        self.assertTrue(access.host_stopped({'status':200,'challenge':True}))
        self.assertTrue(access.host_stopped({'status':None,'error':'connection-error-or-timeout'}))
        self.assertIsNone(access.NoRedirect().redirect_request(None,None,302,'',{},access.BASE+'/file.xlsx'))

    def test_process_lock_is_held_before_ledger_or_transport(self):
        with tempfile.TemporaryDirectory(prefix='opax-sa-lock-') as tmp:
            events=[]
            with patch.object(access,'STAGE',Path(tmp)),patch.object(access.fcntl,'flock',side_effect=lambda *_:events.append('lock')),patch.object(access,'_acquire',side_effect=lambda *_:events.append('request')):
                access.acquire(access.BASE+'/robots.txt','robots')
            self.assertEqual(events,['lock','request'])

    def test_forged_cells_fail_even_when_series_and_hub_are_changed_together(self):
        with tempfile.TemporaryDirectory(prefix='opax-sa-projection-') as tmp:
            stage=Path(tmp)
            for name in ('hub.json','inventory.json','series.jsonl','receipts.json','reconciliation.json'):
                shutil.copyfile(access.STAGE/name,stage/name)
            rows=[json.loads(line) for line in (stage/'series.jsonl').read_text().splitlines()]
            rows[0]['value']=987654321
            hub=json.loads((stage/'hub.json').read_text());hub['rows'][0]['value']=987654321
            (stage/'hub.json').write_text(json.dumps(hub))
            (stage/'series.jsonl').write_text(''.join(json.dumps(r)+'\n' for r in rows))
            with self.assertRaisesRegex(ValueError,'source-cell projection'):
                checked_projection(stage)

    def test_download_needs_source_licence_before_transport(self):
        with tempfile.TemporaryDirectory(prefix='opax-sa-licence-') as tmp:
            root=Path(tmp);cache=root/'cache';cache.mkdir()
            url=access.BASE+'/data/dataset/example/resource/synthetic/download/fixture.xlsx'
            rows=[dict(url=access.BASE+'/copyright',kind='terms',body_retained=True,sha256='terms'),
                  dict(url=access.BASE+'/data/dataset/example/resource/synthetic',kind='resource',body_retained=True,sha256='resource')]
            (cache/'resource').write_text('<a href="'+url+'">download</a>')
            (root/'receipts.json').write_text(json.dumps(rows))
            with patch.object(access,'STAGE',root),patch.object(access,'CACHE',cache),patch.object(access.request,'build_opener') as transport:
                with self.assertRaisesRegex(ValueError,'checked source grant'):access.acquire(url,'download')
                transport.assert_not_called()

    def test_sa_stays_outside_nightly_and_review_gate_denies_unapproved_data(self):
        self.assertNotIn('sa-budget',promoter.DEFAULT_HUBS)
        load=promoter.load
        def pending(path):
            data=load(path)
            return {**data,'status':'pending'} if path==access.STAGE/'review.json' else data
        with patch.object(promoter,'load',side_effect=pending):
            with self.assertRaisesRegex(promoter.PromotionError,'review'):
                promoter.validated_asset('sa-budget')
        expected,_=promoter.validated_asset('sa-budget')
        self.assertEqual((access.ROOT/'portal/public/sa-budget/data.json').read_bytes(),expected)


if __name__=='__main__':unittest.main()
